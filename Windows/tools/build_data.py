# -*- coding: utf-8 -*-
"""
fhlSearch 数据管道
==================
把三个 Excel 题库转换成程序内嵌的紧凑文本资源 bank.txt。

用法:  python tools/build_data.py [--src 目录] [--out 路径] [--check]

题库来源目录按下面的顺序确定（第一个命中的生效）：
  1. --src 命令行参数
  2. 项目根目录下的 bank_src.txt（一行，写题库目录的绝对/相对路径；本机专用，不入库）
  3. 项目根目录下的 data/（仓库自带，默认）

设计要点
--------
* 程序运行时**完全不依赖 Excel/Office**，也不解析 xlsx；数据在编译期被塞进 exe。
* 成语的"单字"列在题库里是 =MID() 公式，这里一律**从成语本身现算**，不读 J-M 列。
* 拼音统一转小写（题库表头写明"无声调小写"）。
* 名士表按"名称所在行为块首"归并合并单元格，每位名士的爱好/社团/描述可能是多条。
* 多值字段之间用 \x1f (Unit Separator) 分隔，字段之间用 \t，记录之间用 \n。
"""
import sys, io, os, argparse, datetime, hashlib
# 输出跟随系统编码（中文 Windows 的 cmd 是 GBK），只把编不出来的字符降级，避免抛异常。
# 需要强制 UTF-8 时，设环境变量 PYTHONIOENCODING=utf-8。
try:
    sys.stdout.reconfigure(errors='replace')
except Exception:
    pass
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))   # 项目根目录
DEFAULT_SRC = os.path.join(ROOT, "data")
DEFAULT_OUT = os.path.join(ROOT, "src", "FhlSearch", "Resources", "bank.txt")
CHECK_OUT = os.path.join(ROOT, "_dev", "bank_preview.txt")
SRC_OVERRIDE = os.path.join(ROOT, "bank_src.txt")

US = "\x1f"      # 多值分隔符
TAB = "\t"
NL = "\n"

# 需求文档里规定的四组菜单选项（顺序按需求原文）
OPT_DYNASTY = ["汉", "明", "秦", "三国", "宋", "唐", "五代十国", "西楚",
               "晋", "先秦", "梦", "南北朝", "清", "元", "辽"]
OPT_HOBBY = ["博文", "乐舞", "品茗", "赏玩", "食珍", "书画", "雅宴", "演武", "弈棋", "妆饰"]
OPT_CLUB = ["百家书院", "千工苑", "金戈馆", "容华社", "棠梨坊", "饕餮居", "五湖商社", "逸趣社", "无"]
OPT_RANK = ["天", "地", "玄"]

warnings = []


def warn(msg):
    warnings.append(msg)
    print("  [警告] " + msg)


def norm(v):
    """单元格 -> 去掉首尾空白的字符串"""
    if v is None:
        return ""
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    return str(v).strip()


def clean_cell(v):
    """单元格 -> 单行文本（换行/制表符/分隔符一律压成空格）"""
    s = norm(v)
    s = s.replace("\r", " ").replace("\n", " ").replace("\t", " ")
    s = s.replace(US, " ")
    while "  " in s:
        s = s.replace("  ", " ")
    return s.strip()


# ---------------------------------------------------------------- 成语
def build_idioms(path):
    print("\n[1/3] 成语题库: " + path)
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = []
    seen = set()
    for i, r in enumerate(ws.iter_rows(min_row=3, values_only=True), start=3):
        idiom = clean_cell(r[0]) if len(r) > 0 else ""
        if not idiom:
            continue
        if len(idiom) != 4:
            warn(f"成语检索 R{i}: '{idiom}' 不是 4 个字，已跳过")
            continue
        if idiom in seen:
            warn(f"成语检索 R{i}: '{idiom}' 重复，已跳过")
            continue
        py = [clean_cell(r[c]).lower() for c in range(1, 5)]
        tn = [clean_cell(r[c]) for c in range(5, 9)]
        ok = True
        for k in range(4):
            if not py[k] or any(not ("a" <= ch <= "z") for ch in py[k]):
                warn(f"成语检索 R{i} '{idiom}': 第{k+1}字拼音 '{py[k]}' 含非小写字母，已跳过该条")
                ok = False
                break
            if tn[k] not in ("1", "2", "3", "4"):
                warn(f"成语检索 R{i} '{idiom}': 第{k+1}字声调 '{tn[k]}' 不是 1-4，已跳过该条")
                ok = False
                break
        if not ok:
            continue
        # 单字从成语现算（不依赖题库里的 =MID() 公式列）
        zi = [idiom[k] for k in range(4)]
        rows.append((idiom, py, tn, zi))
        seen.add(idiom)
    wb.close()
    rows.sort(key=lambda x: (x[1][0], x[0]))
    print(f"  -> {len(rows)} 条成语")
    return rows


# ---------------------------------------------------------------- 诗词
def build_poems(path):
    print("\n[2/3] 诗词题库: " + path)
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    print(f"  工作表: {ws.title}")
    rows = []
    seen = set()
    reps = 0
    for i, r in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        verse = norm(r[0]) if len(r) > 0 else ""
        if not verse:
            continue
        if len(verse) != 5:
            warn(f"诗词检索 R{i}: '{verse}' 不是 5 个字，已跳过")
            continue
        if verse in seen:
            warn(f"诗词检索 R{i}: '{verse}' 重复，已跳过")
            continue
        title = clean_cell(r[1]) if len(r) > 1 else ""
        author = clean_cell(r[2]) if len(r) > 2 else ""
        if len(set(verse)) != 5:
            reps += 1
        rows.append((verse, title, author))
        seen.add(verse)
    wb.close()
    print(f"  -> {len(rows)} 条诗句（其中 {reps} 条含重复字，需要字格里有两个相同字才能拼出）")
    return rows


# ---------------------------------------------------------------- 名士
def build_people(path):
    print("\n[3/3] 名士题库: " + path)
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    print(f"  工作表: {ws.title}")
    rows = list(ws.iter_rows(values_only=True))

    # 找出每个"块"的起始行（A 列有名士名）
    starts = [i for i in range(1, len(rows)) if clean_cell(rows[i][0])]
    blocks = []
    for k, s in enumerate(starts):
        e = starts[k + 1] - 1 if k + 1 < len(starts) else len(rows) - 1
        blocks.append((s, e))

    people = []
    seen = set()
    multi = []
    for s, e in blocks:
        name = clean_cell(rows[s][0])
        if name in seen:
            warn(f"名士检索: '{name}' 重复出现，已跳过后面的块")
            continue
        dyn = clean_cell(rows[s][1])
        rank = clean_cell(rows[s][4])
        job = clean_cell(rows[s][5])
        meet = clean_cell(rows[s][6])
        origin = clean_cell(rows[s][7])
        skill = clean_cell(rows[s][8])

        hobbies, clubs, descs = [], [], []
        for r in range(s, e + 1):
            v = clean_cell(rows[r][2])
            if v and v not in hobbies:
                hobbies.append(v)
            v = clean_cell(rows[r][3])
            if v and v not in clubs:
                clubs.append(v)
            for c in range(9, min(24, len(rows[r]))):
                v = clean_cell(rows[r][c])
                if v:
                    descs.append(v)

        if len(hobbies) > 1:
            multi.append((name, "爱好", hobbies))
        if len(clubs) > 1:
            multi.append((name, "社团", clubs))

        if not dyn:
            warn(f"名士检索: '{name}' 没有朝代")
        if not rank:
            warn(f"名士检索: '{name}' 没有品阶（该名士不会被天/地/玄筛出）")
        if not hobbies:
            warn(f"名士检索: '{name}' 没有爱好")
        if not clubs:
            warn(f"名士检索: '{name}' 没有社团")

        people.append((name, dyn, rank, job, meet, origin, skill, hobbies, clubs, descs))
        seen.add(name)

    wb.close()
    print(f"  -> {len(people)} 位名士")
    print(f"     多位爱好/社团的名士: {len(multi)} 人（如 {multi[0][0] if multi else '-'}）")
    dyns = {}
    for p in people:
        dyns[p[1]] = dyns.get(p[1], 0) + 1
    print(f"     朝代分布: {dict(sorted(dyns.items()))}")
    return people


# ---------------------------------------------------------------- 输出
def compose(idioms, poems, people, src_info):
    L = []
    L.append("FHLS-BANK 1")
    L.append("# 忘川检索 fhlSearch 内嵌题库（本文件由 tools/build_data.py 自动生成，请勿手改）")
    L.append("# 生成时间: " + datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"))
    for k, v in src_info.items():
        L.append("# %s: %s" % (k, v))
    L.append("# 分隔符: 记录=\\n  字段=\\t  多值=\\x1f")
    L.append("")

    def section(tag, lines):
        L.append("[" + tag + "]")
        L.extend(lines)
        L.append("")

    section("OPT_DYNASTY", OPT_DYNASTY)
    section("OPT_HOBBY", OPT_HOBBY)
    section("OPT_CLUB", OPT_CLUB)
    section("OPT_RANK", OPT_RANK)
    section("IDIOM", [TAB.join([i[0]] + i[1] + i[2]) for i in idioms])
    section("POEM", [TAB.join(p) for p in poems])
    section("PERSON", [TAB.join([p[0], p[1], p[2], p[3], p[4], p[5], p[6],
                                 US.join(p[7]), US.join(p[8]), US.join(p[9])])
                       for p in people])
    return NL.join(L) + NL


def resolve_src(cli):
    """题库来源目录: --src > bank_src.txt > data/"""
    if cli:
        return os.path.abspath(cli)
    if os.path.exists(SRC_OVERRIDE):
        with open(SRC_OVERRIDE, encoding="utf-8-sig") as f:
            line = f.readline().strip()
        if line and not line.startswith("#"):
            return os.path.abspath(os.path.join(ROOT, line) if not os.path.isabs(line) else line)
    return DEFAULT_SRC


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--src", default=None, help="题库目录（含三个 xlsx），默认用 data/")
    ap.add_argument("--out", default=DEFAULT_OUT)
    ap.add_argument("--check", action="store_true", help="只生成到 _dev 做检查，不覆盖正式资源")
    args = ap.parse_args()

    out = CHECK_OUT if args.check else args.out
    src_dir = resolve_src(args.src)

    f_idiom = os.path.join(src_dir, "成语检索.xlsx")
    f_poem = os.path.join(src_dir, "诗词检索.xlsx")
    f_person = os.path.join(src_dir, "名士检索.xlsx")
    for f in (f_idiom, f_poem, f_person):
        if not os.path.exists(f):
            print("找不到题库文件: " + f)
            print()
            print("题库目录当前取的是: " + src_dir)
            print("请把三个 xlsx（成语检索 / 诗词检索 / 名士检索）放进该目录，")
            print("或者用 --src <目录> 指定别处，也可以在项目根目录建一个 bank_src.txt")
            print("写上一行题库目录的路径。")
            return 2

    print("=" * 78)
    print("忘川检索 数据管道")
    print("=" * 78)
    print("题库目录: " + src_dir)

    idioms = build_idioms(f_idiom)
    poems = build_poems(f_poem)
    people = build_people(f_person)

    src_info = {}
    for tag, f in (("成语题库", f_idiom), ("诗词题库", f_poem), ("名士题库", f_person)):
        st = os.stat(f)
        src_info[tag] = "%s  (%d bytes, 修改于 %s)" % (
            os.path.basename(f), st.st_size,
            datetime.datetime.fromtimestamp(st.st_mtime).strftime("%Y-%m-%d %H:%M"))

    text = compose(idioms, poems, people, src_info)

    os.makedirs(os.path.dirname(out), exist_ok=True)
    with open(out, "w", encoding="utf-8", newline="\n") as f:
        f.write(text)

    size = os.path.getsize(out)
    sha = hashlib.sha256(text.encode("utf-8")).hexdigest()[:16]
    print("\n" + "=" * 78)
    print(f"输出: {out}")
    print(f"大小: {size} bytes   sha256[0:16]: {sha}")
    print(f"统计: 成语 {len(idioms)} 条 | 诗句 {len(poems)} 条 | 名士 {len(people)} 位")
    print(f"警告: {len(warnings)} 条")
    print("=" * 78)
    return 0


if __name__ == "__main__":
    sys.exit(main())
